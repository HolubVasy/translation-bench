"""Convert source articles (.pdf or .docx) to segmented JSON.

Segment ids are structural addresses: S2-p3 means section 2, paragraph 3.
Headings are S2-h (or S2-h2 for a second heading inside the section),
images are S2-img1. A new section starts at every level-1 heading.
Images become [IMAGE_n] placeholders. For PDF input the image regions are
also saved to data/source/images/a{n}/IMAGE_{n}.png so the winning
translation can be assembled back into a document (src.assemble).

PDF has no styles, so structure is recovered from layout:
- body font size = character-weighted median of all text blocks;
- a short block clearly larger than body is a heading, the largest
  heading size is level 1, any smaller heading size is level 2;
- page furniture is dropped: bare page numbers in the margin, and short
  text repeated on many pages anywhere (running heads, download lines);
- vertical text (side watermarks) is ignored;
- layout debris is dropped: blocks with too few words or letters to be
  prose (axis labels, legends, formula fragments, number grids);
- the article ends at a back-matter heading (References by default);
- lines are joined with spaces, words hyphenated across lines are glued.
Settings live under `convert` in config.yaml. Review data/source/a{n}.json
once before a paid run.

Usage: python -m src.convert articles/1.pdf [more files ...]
Output: data/source/a{n}.json (n from the digits in the file name,
otherwise sequential).
"""

import json
import re
import sys
from collections import Counter
from pathlib import Path

SRC_DIR = Path("data/source")
IMG_ROOT = SRC_DIR / "images"

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
DRAWING_TAGS = (W_NS + "drawing", W_NS + "object", W_NS + "pict")

MARGIN_FRACTION = 0.08
HEADING_RATIO_MIN = 1.15
HEADING_MAX_CHARS = 200
MIN_IMAGE_SIDE = 50
SOFT_HYPHEN = "\u00ad"
INVISIBLE_RE = re.compile("[\u200b\u200c\u200d\u2060\ufeff]")
URL_BREAK_RE = re.compile(r"(https?://\S*[./-])\s+(?=[\w/]\S*[./])")  # URL broken across lines
FURNITURE_MAX_CHARS = 150
MIN_PARAGRAPH_WORDS = 3
MIN_LETTER_RATIO = 0.55
DEFAULT_STOP_AT = ("references", "bibliography", "literature cited")


# ---------------------------------------------------------------- segments


def build_segments(items):
    """items: dicts with kind in (heading, para, image), level for headings.
    Returns the flat segment list with structural ids."""
    segments = []
    section = 0
    p_count = h_count = img_in_section = img_total = 0
    for item in items:
        kind = item["kind"]
        level = item.get("level", 0)
        if (kind == "heading" and level == 1) or section == 0:
            section += 1
            p_count = h_count = img_in_section = 0
        if kind == "image":
            img_total += 1
            img_in_section += 1
            segments.append(
                {
                    "id": "S%d-img%d" % (section, img_in_section),
                    "section": section,
                    "type": "image",
                    "text": "[IMAGE_%d]" % img_total,
                }
            )
            item["image_number"] = img_total
        elif kind == "heading":
            h_count += 1
            seg_id = "S%d-h" % section if h_count == 1 else "S%d-h%d" % (section, h_count)
            segments.append(
                {"id": seg_id, "section": section, "type": "h%d" % level, "text": item["text"]}
            )
        else:
            p_count += 1
            segments.append(
                {"id": "S%d-p%d" % (section, p_count), "section": section,
                 "type": "p", "text": item["text"]}
            )
    return segments


# ---------------------------------------------------------------- docx


def _para_has_image(paragraph):
    for run in paragraph.runs:
        for tag in DRAWING_TAGS:
            if run._element.find(tag) is not None:
                return True
    return False


def _docx_heading_level(paragraph):
    name = (paragraph.style.name or "").lower()
    match = re.match(r"heading\s+(\d+)", name)
    if match:
        return int(match.group(1))
    if name == "title":
        return 1
    return 0


def docx_items(path):
    from docx import Document

    items = []
    for paragraph in Document(str(path)).paragraphs:
        text = paragraph.text.strip()
        has_image = _para_has_image(paragraph)
        if not text and not has_image:
            continue
        level = _docx_heading_level(paragraph)
        if has_image:
            items.append({"kind": "image"})
            if not text:
                continue
        if level > 0:
            items.append({"kind": "heading", "level": level, "text": text})
        else:
            items.append({"kind": "para", "text": text})
    return items


# ---------------------------------------------------------------- pdf


def join_lines(lines):
    text = ""
    for line in lines:
        line = INVISIBLE_RE.sub("", line.replace(SOFT_HYPHEN, "")).strip()
        if not line:
            continue
        if not text:
            text = line
        elif text.endswith("-") and line[:1].islower():
            text = text[:-1] + line
        else:
            text = text + " " + line
    text = re.sub(r"\s+", " ", text).strip()
    return URL_BREAK_RE.sub(r"\1", text)


def _weighted_median(pairs):
    """pairs: (value, weight)."""
    pairs = sorted(pairs)
    total = sum(w for _, w in pairs)
    acc = 0
    for value, weight in pairs:
        acc += weight
        if acc >= total / 2:
            return value
    return pairs[-1][0] if pairs else 0


def _norm_repeat(text):
    return re.sub(r"\d+", "#", text).strip().lower()


def _is_horizontal(line):
    dx, dy = line.get("dir", (1, 0))
    return abs(dy) < 0.1 and dx > 0


def _block_text(block):
    """Joined text and font sizes of the horizontal lines of a text block."""
    lines, sizes = [], []
    for line in block.get("lines", []):
        if not _is_horizontal(line):
            continue
        text = "".join(span["text"] for span in line["spans"])
        if text.strip():
            lines.append(text)
            sizes += [span["size"] for span in line["spans"] if span["text"].strip()]
    return join_lines(lines), sizes


def pdf_raw_blocks(path):
    import pymupdf

    doc = pymupdf.open(str(path))
    blocks = []
    for pno, page in enumerate(doc):
        height = page.rect.height
        for block in page.get_text("dict")["blocks"]:
            x0, y0, x1, y1 = block["bbox"]
            in_margin = y1 < height * MARGIN_FRACTION or y0 > height * (1 - MARGIN_FRACTION)
            if block["type"] == 1:
                if (x1 - x0) >= MIN_IMAGE_SIDE and (y1 - y0) >= MIN_IMAGE_SIDE:
                    blocks.append({"page": pno, "kind": "image", "bbox": block["bbox"]})
                continue
            text, sizes = _block_text(block)
            if text:
                blocks.append({"page": pno, "kind": "text", "text": text,
                               "size": max(sizes), "margin": in_margin})
    return doc, blocks


# ---- noise filters: each takes (block, context) and says whether to drop it


def is_page_furniture(block, ctx):
    """Running headers, footers and page numbers: bare numbers in the margin,
    or short text repeated on many pages anywhere (watermarks, download lines)."""
    if block["margin"] and re.fullmatch(r"[\d\s\-/|.]+", block["text"]):
        return True
    if ctx["n_pages"] < 2 or len(block["text"]) > FURNITURE_MAX_CHARS:
        return False
    return ctx["repeats"][_norm_repeat(block["text"])] >= ctx["repeat_min"]


def _real_words(text):
    """Words with at least three letters; formula symbols and indices are not words."""
    return [w for w in text.split() if sum(ch.isalpha() for ch in w) >= 3]


def is_layout_debris(block, ctx):
    """Axis labels, legend entries, formula fragments and number grids:
    too few real words or too few letters to be prose. A heading needs
    only one real word."""
    text = block["text"]
    letters = sum(ch.isalpha() for ch in text)
    if letters < MIN_LETTER_RATIO * len(text):
        return True
    needed = 1 if ctx["is_heading"](block) else MIN_PARAGRAPH_WORDS
    return len(_real_words(text)) < needed


NOISE_FILTERS = {
    "page_furniture": is_page_furniture,
    "layout_debris": is_layout_debris,
}


def is_stop_heading(text, stop_at):
    """True for a back-matter heading such as References: the article ends there."""
    cleaned = re.sub(r"^[^A-Za-z]+", "", text).strip().rstrip(":").lower()
    return cleaned in stop_at


def _noise_context(doc, blocks):
    repeats = Counter(_norm_repeat(b["text"]) for b in blocks if b["kind"] == "text")
    return {"n_pages": len(doc), "repeats": repeats, "repeat_min": max(2, int(len(doc) * 0.3))}


def _heading_classifier(texts):
    body = _weighted_median([(round(b["size"], 1), len(b["text"])) for b in texts]) or 10
    heading_sizes = sorted(
        {round(b["size"], 1) for b in texts
         if b["size"] >= body * HEADING_RATIO_MIN and len(b["text"]) <= HEADING_MAX_CHARS},
        reverse=True,
    )

    def level_of(block):
        size = round(block["size"], 1)
        if size not in heading_sizes or len(block["text"]) > HEADING_MAX_CHARS:
            return 0
        return 1 if size == heading_sizes[0] else 2

    return level_of


def pdf_items(path, settings=None):
    settings = settings or {}
    stop_at = {s.lower() for s in settings.get("stop_at", DEFAULT_STOP_AT)}
    doc, blocks = pdf_raw_blocks(path)
    ctx = _noise_context(doc, blocks)
    texts = [b for b in blocks if b["kind"] == "text" and not is_page_furniture(b, ctx)]
    level_of = _heading_classifier(texts)
    ctx["is_heading"] = lambda b: level_of(b) > 0
    items, image_blocks = [], []
    for block in blocks:
        if block["kind"] == "image":
            items.append({"kind": "image"})
            image_blocks.append((items[-1], block))
            continue
        if is_stop_heading(block["text"], stop_at):
            break
        if any(drop(block, ctx) for drop in NOISE_FILTERS.values()):
            continue
        level = level_of(block)
        kind = "heading" if level else "para"
        items.append({"kind": kind, "level": level, "text": block["text"]} if level
                     else {"kind": kind, "text": block["text"]})
    return doc, items, image_blocks


def save_pdf_images(doc, image_blocks, img_dir):
    import pymupdf

    img_dir.mkdir(parents=True, exist_ok=True)
    for old in img_dir.glob("IMAGE_*"):
        old.unlink()
    for item, block in image_blocks:
        page = doc[block["page"]]
        pix = page.get_pixmap(clip=pymupdf.Rect(block["bbox"]), dpi=150)
        pix.save(str(img_dir / ("IMAGE_%d.png" % item["image_number"])))


# ---------------------------------------------------------------- main


def load_settings():
    path = Path("config.yaml")
    if not path.exists():
        return {}
    import yaml

    return (yaml.safe_load(path.read_text(encoding="utf-8")) or {}).get("convert") or {}


def convert_file(path, article_id, settings=None):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        img_dir = IMG_ROOT / ("a%s" % article_id)
        doc, items, image_blocks = pdf_items(path, settings)
        segments = build_segments(items)
        save_pdf_images(doc, image_blocks, img_dir)
    elif suffix == ".docx":
        segments = build_segments(docx_items(path))
    else:
        raise ValueError("unsupported input %s (use .pdf or .docx)" % path.name)
    return {"article": str(article_id), "segments": segments}


def main(argv):
    if not argv:
        print(__doc__)
        return 1
    SRC_DIR.mkdir(parents=True, exist_ok=True)
    settings = load_settings()
    next_seq = 1
    for raw in [a for a in argv if not a.startswith("--")]:
        path = Path(raw)
        digits = re.sub(r"\D", "", path.stem)
        article_id = int(digits) if digits else next_seq
        next_seq = article_id + 1
        data = convert_file(path, article_id, settings)
        out = SRC_DIR / ("a%s.json" % article_id)
        out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        segs = data["segments"]
        print(
            "%s -> %s (%d segments, %d sections, %d images, %d words)"
            % (path.name, out, len(segs), segs[-1]["section"] if segs else 0,
               sum(1 for s in segs if s["type"] == "image"),
               sum(len(s["text"].split()) for s in segs if s["type"] != "image"))
        )
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
