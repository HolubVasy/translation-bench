"""Assemble a translation back into a .docx document.

Headings and paragraphs come from the translated segments; [IMAGE_n]
placeholders are replaced by the images saved at convert time
(data/source/images/a{n}/IMAGE_{n}.png, PDF input only). A placeholder
without an image file stays as text.

The translator defaults to the overall winner by Borda points (judges that
failed calibration excluded); pass --translator to choose another one.

Usage: python -m src.assemble [--translator M4] [a1 a2 ...]
Output: data/report/assembled__a{n}__M{k}.docx
"""

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.shared import Inches

SRC_DIR = Path("data/source")
TR_DIR = Path("data/translations")
JUDG_DIR = Path("data/judgments")
CAL_DIR = Path("data/calibration")
OUT_DIR = Path("data/report")


def failed_judges():
    failed = set()
    for path in CAL_DIR.glob("cal__by__*.json"):
        record = json.loads(path.read_text(encoding="utf-8"))
        if not record.get("passed", True):
            failed.add(record["judge"])
    return failed


def borda_winner():
    failed = failed_judges()
    all_judges = {json.loads(p.read_text(encoding="utf-8"))["judge"]
                  for p in JUDG_DIR.glob("ev__*.json")}
    # Same rule as aggregate.py: exclude failed judges, unless none would remain.
    failed = failed & all_judges
    if failed == all_judges:
        failed = set()
    points = defaultdict(list)
    for path in JUDG_DIR.glob("ev__*.json"):
        judgment = json.loads(path.read_text(encoding="utf-8"))
        if judgment["judge"] in failed:
            continue
        ranking = judgment["ranking"]
        for pos, code in enumerate(ranking):
            points[code].append(len(ranking) - pos)
    if not points:
        return None
    return max(points, key=lambda c: sum(points[c]) / len(points[c]))


def assemble(article, code):
    tr_path = TR_DIR / ("tr__%s__%s.json" % (article, code))
    if not tr_path.exists():
        print("missing %s" % tr_path)
        return None
    translation = json.loads(tr_path.read_text(encoding="utf-8"))
    img_dir = SRC_DIR / "images" / article
    doc = Document()
    for seg in translation["segments"]:
        kind, text = seg["type"], seg["text"]
        if kind == "image":
            match = re.search(r"\[IMAGE_(\d+)\]", text)
            img = img_dir / ("IMAGE_%s.png" % match.group(1)) if match else None
            if img and img.exists():
                doc.add_picture(str(img), width=Inches(6))
            else:
                doc.add_paragraph(text)
        elif kind.startswith("h"):
            level = int(kind[1:]) if kind[1:].isdigit() else 1
            doc.add_heading(text, level=min(level, 4))
        else:
            doc.add_paragraph(text)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / ("assembled__%s__%s.docx" % (article, code))
    doc.save(str(out))
    return out


def main(argv):
    parser = argparse.ArgumentParser()
    parser.add_argument("--translator")
    parser.add_argument("articles", nargs="*")
    args = parser.parse_args(argv)
    code = args.translator or borda_winner()
    if not code:
        print("no judgments yet; pass --translator explicitly")
        return 1
    articles = args.articles or sorted(p.stem for p in SRC_DIR.glob("a*.json"))
    built = 0
    for article in articles:
        out = assemble(article, code)
        if out:
            built += 1
            print("assembled %s" % out)
    return 0 if built else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
